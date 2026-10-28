from pathlib import Path
from datetime import datetime
import sys, html, traceback, warnings, webbrowser

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
import plotly.graph_objects as go
import plotly.io as pio
from pptx import Presentation
from pptx.util import Inches, Pt
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
try:
    import webview
except ImportError:
    webview = None

APP_NAME = "배터리 품질 분석 자동화"
BASE_DIR = Path(sys.executable).resolve().parent if getattr(sys, "frozen", False) else Path(__file__).resolve().parent
RAW_DIR = BASE_DIR / "Raw_Data"
RESULT_DIR = BASE_DIR / "Results"
FIGURE_DIR = RESULT_DIR / "figures"
for folder in (RAW_DIR, RESULT_DIR, FIGURE_DIR):
    folder.mkdir(parents=True, exist_ok=True)

HTML_OUTPUT = RESULT_DIR / "battery_quality_dashboard.html"
PPT_OUTPUT = RESULT_DIR / "battery_quality_report.pptx"
LOG_OUTPUT = RESULT_DIR / "run_log.txt"
RESULT_CSV = RESULT_DIR / "battery_quality_judgment_results.csv"
SPEC_FILE = RAW_DIR / "specifications.xlsx"
SUPPORTED_EXTENSIONS = {".csv", ".xlsx", ".xls"}

COLUMN_ALIASES = {
    "sample_id": ["sample_id", "sampleid", "id", "시료번호", "샘플번호", "셀번호", "배터리번호"],
    "capacity_ah": ["capacity_ah", "capacity", "용량", "방전용량", "용량(ah)", "용량 (ah)", "방전용량(ah)"],
    "voltage_v": ["voltage_v", "voltage", "전압", "전압(v)", "전압 (v)"],
    "ir_mohm": ["ir_mohm", "ir", "resistance", "저항", "내부저항", "저항(mohm)", "내부저항(mohm)", "저항 (mΩ)", "내부저항 (mΩ)"],
    "thickness_mm": ["thickness_mm", "thickness", "두께", "두께(mm)", "두께 (mm)"],
    "weight_g": ["weight_g", "weight", "중량", "무게", "중량(g)", "무게(g)", "중량 (g)"],
    "quality_label": ["quality_label", "quality", "판정", "품질판정", "검사결과", "result", "ok_ng", "pass_fail"],
    "process": ["process", "공정", "공정명", "공정단계", "process_name"],
}
DISPLAY_NAMES = {
    "capacity_ah": "용량 (Ah)", "voltage_v": "전압 (V)", "ir_mohm": "내부저항 (mΩ)",
    "thickness_mm": "두께 (mm)", "weight_g": "중량 (g)", "quality_label": "전체 품질 판정",
}
NUMERIC_COLUMNS = ["capacity_ah", "voltage_v", "ir_mohm", "thickness_mm", "weight_g"]
JUDGMENT_LABELS = ("OK", "NG", "미측정", "미판정")


def write_log(message):
    stamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    try:
        with LOG_OUTPUT.open("a", encoding="utf-8") as f:
            f.write(f"[{stamp}] {message}\n")
    except OSError:
        pass
    print(f"[{stamp}] {message}")


def initialize_log():
    LOG_OUTPUT.write_text(f"{APP_NAME}\n실행 시간: {datetime.now()}\n{'=' * 60}\n", encoding="utf-8")


def find_input_file():
    files = [p for p in RAW_DIR.iterdir() if p.is_file() and p.suffix.lower() in SUPPORTED_EXTENSIONS
             and p.name.lower() != SPEC_FILE.name.lower() and not p.name.startswith("~$")
             and not p.name.startswith("battery_quality_judgment_results")]
    if not files:
        raise FileNotFoundError(f"분석할 데이터 파일이 없습니다.\n{RAW_DIR} 폴더에 CSV 또는 Excel 파일을 넣어 주세요.")
    if len(files) > 1:
        raise ValueError("Raw_Data 폴더에는 측정 데이터 파일을 한 번에 하나만 넣어 주세요.\n" +
                         "\n".join(f"- {p.name}" for p in files))
    return files[0]


def read_input_file(path):
    if path.suffix.lower() == ".csv":
        last_error = None
        for encoding in ("utf-8-sig", "cp949", "euc-kr"):
            try:
                return pd.read_csv(path, encoding=encoding)
            except UnicodeDecodeError as exc:
                last_error = exc
        raise ValueError("CSV 문자 인코딩을 확인할 수 없습니다.") from last_error
    if path.suffix.lower() == ".xls":
        return pd.read_excel(path, engine="xlrd")
    return pd.read_excel(path, engine="openpyxl")


def standardize_columns(df):
    def norm(x):
        return str(x).strip().lower().replace(" ", "").replace("_", "").replace("-", "")
    aliases = {key: {norm(x) for x in values} for key, values in COLUMN_ALIASES.items()}
    rename = {}
    for col in df.columns:
        for key, values in aliases.items():
            if norm(col) in values:
                rename[col] = key
                break
    df = df.rename(columns=rename)
    if df.columns.duplicated().any():
        raise ValueError(f"표준화 후 중복 컬럼이 발생했습니다: {df.columns[df.columns.duplicated()].tolist()}")
    return df


def get_numeric_columns(df):
    # 자동 생성된 판정 컬럼은 제외하고 원래의 측정 항목만 반환
    return [c for c in NUMERIC_COLUMNS if c in df.columns and pd.to_numeric(df[c], errors="coerce").notna().any()]


def normalize_quality_label(series):
    labels = series.astype("string").str.strip().str.upper()
    return labels.replace({
        "합격": "OK", "양품": "OK", "PASS": "OK", "P": "OK", "정상": "OK", "GOOD": "OK",
        "불합격": "NG", "불량": "NG", "FAIL": "NG", "F": "NG", "이상": "NG", "BAD": "NG",
    })


def prepare_data(df):
    if df.empty:
        raise ValueError("데이터 파일에 분석할 행이 없습니다.")
    df = standardize_columns(df)
    for col in NUMERIC_COLUMNS:
        if col in df.columns:
            # 빈칸과 숫자가 아닌 문자는 결측값으로 처리
            df[col] = pd.to_numeric(df[col].replace(r"^\s*$", np.nan, regex=True), errors="coerce")
    if "quality_label" in df.columns:
        # 규격 기반 판정과 비교할 수 있도록 원본 판정을 보존
        df["quality_label_original"] = normalize_quality_label(df["quality_label"])
        df = df.drop(columns=["quality_label"])
    df = df.dropna(how="all", subset=[c for c in df.columns if c != "quality_label_original"]).reset_index(drop=True)
    if not get_numeric_columns(df):
        raise ValueError("분석 가능한 숫자형 품질 컬럼이 없습니다. 용량, 전압, 내부저항, 두께, 중량 컬럼을 확인하세요.")
    return df


def style_spec_sheet(ws):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1D4ED8")
        cell.alignment = Alignment(horizontal="center")
    for col, width in {"A": 24, "B": 22, "C": 16, "D": 16, "E": 16}.items():
        ws.column_dimensions[col].width = width
    for row in range(2, ws.max_row + 1):
        for col in (3, 4, 5):
            ws.cell(row, col).number_format = "0.########"
            ws.cell(row, col).fill = PatternFill("solid", fgColor="FFF7D6")


def ensure_spec_template(df):
    keys = get_numeric_columns(df)
    if SPEC_FILE.exists():
        wb = load_workbook(SPEC_FILE)
        ws = wb.active
        headers = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
        expected = ["분석항목(표준명)", "분석항목", "Target", "LSL", "USL"]
        if headers[:5] != expected:
            raise ValueError(f"{SPEC_FILE.name} 첫 행의 열 이름을 다음 순서로 맞춰 주세요: {', '.join(expected)}")
        existing = {str(ws.cell(r, 1).value).strip() for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value}
        for key in keys:
            if key not in existing:
                ws.append([key, DISPLAY_NAMES.get(key, key), None, None, None])
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "규격 설정"
        ws.append(["분석항목(표준명)", "분석항목", "Target", "LSL", "USL"])
        for key in keys:
            ws.append([key, DISPLAY_NAMES.get(key, key), None, None, None])
    style_spec_sheet(ws)
    wb.save(SPEC_FILE)


def optional_number(value, row, field):
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        raise ValueError(f"규격 설정 파일 {row}행의 {field}는 숫자로 입력해야 합니다.")
    if not np.isfinite(number):
        raise ValueError(f"규격 설정 파일 {row}행의 {field}는 유한한 숫자여야 합니다.")
    return number


def load_specifications(df):
    ensure_spec_template(df)
    ws = load_workbook(SPEC_FILE, data_only=True).active
    specs = {}
    for r in range(2, ws.max_row + 1):
        key = ws.cell(r, 1).value
        if not key:
            continue
        key = str(key).strip()
        if key not in NUMERIC_COLUMNS:
            continue
        spec = {"target": optional_number(ws.cell(r, 3).value, r, "Target"),
                "lsl": optional_number(ws.cell(r, 4).value, r, "LSL"),
                "usl": optional_number(ws.cell(r, 5).value, r, "USL")}
        if spec["lsl"] is not None and spec["usl"] is not None and spec["lsl"] >= spec["usl"]:
            raise ValueError(f"규격 설정 파일 {r}행({DISPLAY_NAMES.get(key,key)}): LSL은 USL보다 작아야 합니다.")
        if spec["target"] is not None and spec["lsl"] is not None and spec["target"] < spec["lsl"]:
            write_log(f"주의: {key} Target이 LSL보다 작습니다.")
        if spec["target"] is not None and spec["usl"] is not None and spec["target"] > spec["usl"]:
            write_log(f"주의: {key} Target이 USL보다 큽니다.")
        specs[key] = spec
    for key in get_numeric_columns(df):
        specs.setdefault(key, {"target": None, "lsl": None, "usl": None})
    return specs


def apply_spec_judgments(df, specs):
    """각 측정값을 규격과 비교한다. 전체 NG는 하나라도 NG, 전체 OK는 모든 항목 OK일 때만 부여."""
    result = df.copy()
    measured_cols = get_numeric_columns(result)
    judgment_columns = []
    for key in measured_cols:
        spec = specs.get(key, {})
        lsl, usl = spec.get("lsl"), spec.get("usl")
        judgment_col = f"{key}_판정"
        judgment_columns.append(judgment_col)
        values = pd.to_numeric(result[key], errors="coerce")
        labels = pd.Series("미측정", index=result.index, dtype="object")
        if lsl is None and usl is None:
            labels.loc[values.notna()] = "미판정"
        else:
            valid = values.notna()
            passed = pd.Series(True, index=result.index)
            if lsl is not None:
                passed &= values >= lsl
            if usl is not None:
                passed &= values <= usl
            labels.loc[valid & passed] = "OK"
            labels.loc[valid & ~passed] = "NG"
        result[judgment_col] = labels

    if not judgment_columns:
        result["quality_label"] = "미판정"
        return result
    statuses = result[judgment_columns]
    has_ng = statuses.eq("NG").any(axis=1)
    all_ok = statuses.eq("OK").all(axis=1)
    result["quality_label"] = "미판정"
    result.loc[all_ok, "quality_label"] = "OK"
    result.loc[has_ng, "quality_label"] = "NG"
    return result


def calculate_statistics(df):
    cols = get_numeric_columns(df)
    stats = df[cols].describe().T.rename(columns={
        "count": "데이터 수", "mean": "평균", "std": "표준편차", "min": "최솟값",
        "25%": "1사분위수", "50%": "중앙값", "75%": "3사분위수", "max": "최댓값",
    })
    stats.index = [DISPLAY_NAMES.get(c, c) for c in stats.index]
    return stats.round(4)


def calculate_quality_summary(df):
    labels = df.get("quality_label", pd.Series("미판정", index=df.index)).fillna("미판정")
    ok, ng = int((labels == "OK").sum()), int((labels == "NG").sum())
    unclassified = int((~labels.isin(["OK", "NG"])).sum())
    classified = ok + ng
    return {"total": len(df), "ok": ok, "ng": ng, "unclassified": unclassified,
            "classified": classified, "yield_rate": ok / classified * 100 if classified else None}


def calculate_process_summary(df):
    """공정 컬럼이 있을 때 공정별 전체 OK/NG를 비교한다."""
    if "process" not in df.columns:
        return pd.DataFrame()
    work = df.copy()
    work["process"] = work["process"].astype("string").str.strip()
    work.loc[work["process"].isna() | (work["process"] == ""), "process"] = "공정 미입력"
    rows = []
    for process, group in work.groupby("process", dropna=False):
        labels = group["quality_label"].fillna("미판정")
        ok, ng = int((labels == "OK").sum()), int((labels == "NG").sum())
        classified = ok + ng
        rows.append({"공정": str(process), "전체 수": len(group), "OK 수": ok, "NG 수": ng,
                     "미판정/미측정 수": len(group) - classified,
                     "합격률(OK/(OK+NG))": ok / classified * 100 if classified else np.nan})
    return pd.DataFrame(rows).sort_values("공정").reset_index(drop=True)


def calculate_item_summary(df):
    rows = []
    for key in get_numeric_columns(df):
        col = f"{key}_판정"
        labels = df[col] if col in df.columns else pd.Series("미판정", index=df.index)
        ok = int((labels == "OK").sum())
        ng = int((labels == "NG").sum())
        missing = int((labels == "미측정").sum())
        unjudged = int((labels == "미판정").sum())
        judged = ok + ng
        rows.append({"품질 항목": DISPLAY_NAMES.get(key, key), "OK 수": ok, "NG 수": ng,
                     "미측정 수": missing, "미판정 수": unjudged,
                     "판정 가능 수": judged, "합격률(OK/(OK+NG))": ok / judged * 100 if judged else np.nan,
                     "불량률(NG/(OK+NG))": ng / judged * 100 if judged else np.nan})
    return pd.DataFrame(rows)


def calculate_capability(df, specs):
    rows = []
    for key in get_numeric_columns(df):
        values = df[key].dropna().astype(float)
        spec = specs.get(key, {})
        mean = float(values.mean()) if len(values) else np.nan
        std = float(values.std(ddof=1)) if len(values) >= 2 else np.nan
        lsl, usl, target = spec.get("lsl"), spec.get("usl"), spec.get("target")
        cp = cpu = cpl = cpk = None
        if np.isfinite(std) and std > 0:
            if lsl is not None and usl is not None:
                cp = (usl - lsl) / (6 * std)
                cpu, cpl = (usl - mean) / (3 * std), (mean - lsl) / (3 * std)
                cpk = min(cpu, cpl)
            else:
                if usl is not None: cpu = (usl - mean) / (3 * std)
                if lsl is not None: cpl = (mean - lsl) / (3 * std)
        delta = mean - target if target is not None and np.isfinite(mean) else None
        rows.append({"표준명": key, "품질 항목": DISPLAY_NAMES.get(key, key), "데이터 수": len(values),
                     "Target": target, "LSL": lsl, "USL": usl, "평균": mean, "표준편차": std,
                     "평균-Target": delta, "절대편차": abs(delta) if delta is not None else None,
                     "Cp": cp, "Cpu": cpu, "Cpl": cpl, "Cpk": cpk})
    return pd.DataFrame(rows)


def fmt(x, digits=4):
    return "—" if x is None or pd.isna(x) else f"{x:,.{digits}f}"


def table_html(df, float_digits=2):
    return df.to_html(index=False, border=0, classes="stats-table", na_rep="—",
                      float_format=lambda x: f"{x:.{float_digits}f}", escape=True)


def build_plot_sections(df, specs):
    sections, figures = [], []
    cols = get_numeric_columns(df)
    for col in cols:
        vals = df[col].dropna()
        fig = go.Figure(data=[go.Histogram(x=vals, nbinsx=min(30, max(5, int(np.sqrt(len(vals))*2))))])
        sp = specs.get(col, {})
        for field, label, color in (("lsl", "LSL", "#dc2626"), ("target", "Target", "#2563eb"), ("usl", "USL", "#dc2626")):
            if sp.get(field) is not None:
                fig.add_vline(x=sp[field], line_dash="dash", line_color=color, annotation_text=f"{label}: {sp[field]:g}")
        fig.update_layout(title=f"{DISPLAY_NAMES[col]} 분포", xaxis_title=DISPLAY_NAMES[col], yaxis_title="데이터 수", template="plotly_white", height=360)
        figures.append((f"{DISPLAY_NAMES[col]} 분포", fig))

        judgment_col = f"{col}_판정"
        if judgment_col in df.columns:
            box = go.Figure()
            for label in ("OK", "NG"):
                group = df.loc[df[judgment_col] == label, col].dropna()
                if not group.empty:
                    box.add_trace(go.Box(y=group, name=label, boxpoints="outliers"))
            if box.data:
                box.update_layout(title=f"{DISPLAY_NAMES[col]} 항목 자체의 OK/NG 비교", xaxis_title="해당 항목 판정", yaxis_title=DISPLAY_NAMES[col], template="plotly_white", height=360)
                figures.append((f"{DISPLAY_NAMES[col]} 자체 OK/NG 비교", box))

    if len(cols) >= 2:
        corr = df[cols].corr()
        fig = go.Figure(data=[go.Heatmap(z=corr.values, x=[DISPLAY_NAMES[c] for c in corr.columns],
            y=[DISPLAY_NAMES[c] for c in corr.index], zmin=-1, zmax=1, zmid=0, colorscale="RdBu",
            text=np.round(corr.values, 2), texttemplate="%{text}", colorbar={"title": "상관계수"})])
        fig.update_layout(title="품질 항목 간 상관관계", template="plotly_white", height=430)
        figures.append(("품질 항목 간 상관관계", fig))

    # 전체 판정이 NG인 셀과 OK인 셀의 다른 측정 항목 분포 비교
    if "quality_label" in df.columns:
        box = go.Figure()
        for label in ("OK", "NG"):
            vals = df.loc[df["quality_label"] == label, "capacity_ah"].dropna() if "capacity_ah" in df.columns else pd.Series(dtype=float)
            if not vals.empty:
                box.add_trace(go.Box(y=vals, name=label, boxpoints="outliers"))
        if box.data:
            box.update_layout(title="전체 OK/NG 셀의 용량 비교", xaxis_title="전체 품질 판정", yaxis_title="용량 (Ah)", template="plotly_white", height=360)
            figures.append(("전체 OK/NG 셀 용량 비교", box))

    for i, (title, fig) in enumerate(figures):
        plot = pio.to_html(fig, full_html=False, include_plotlyjs=True if i == 0 else False,
                           config={"responsive": True, "displaylogo": False, "scrollZoom": True})
        sections.append('<section class="panel"><h2>'+html.escape(title)+'</h2>'+plot+'</section>')
    return sections


def build_dashboard(df, stats, summary, source_name, specs, capability, item_summary, process_summary):
    stats_html = stats.to_html(border=0, classes="stats-table", na_rep="—", float_format=lambda x: f"{x:.4f}")
    cap_cols = ["품질 항목", "데이터 수", "Target", "LSL", "USL", "평균", "표준편차", "평균-Target", "절대편차", "Cp", "Cpu", "Cpl", "Cpk"]
    cap = capability[cap_cols].copy()
    for col in cap_cols:
        if col not in ("품질 항목", "데이터 수"):
            cap[col] = cap[col].map(fmt)
    cap_html = cap.to_html(index=False, border=0, classes="stats-table", escape=True)
    item_html = table_html(item_summary, 2)
    if process_summary.empty:
        process_html = '<p class="note">입력 데이터에 공정 컬럼이 없어 공정별 비교를 생략했습니다. 공정별 분석이 필요하면 측정 데이터에 process 또는 공정 컬럼을 추가하세요.</p>'
    else:
        proc_view = process_summary.copy()
        proc_view["합격률(OK/(OK+NG))"] = proc_view["합격률(OK/(OK+NG))"].map(lambda x: "—" if pd.isna(x) else f"{x:.2f}%")
        process_html = proc_view.to_html(index=False, border=0, classes="stats-table", escape=True)
    yield_text = "N/A" if summary["yield_rate"] is None else f'{summary["yield_rate"]:.2f}%'
    ok_text, ng_text = f'{summary["ok"]:,}', f'{summary["ng"]:,}'
    unclassified_text = f'{summary["unclassified"]:,}'
    graphs = "\n".join(build_plot_sections(df, specs))
    page = f"""<!doctype html><html lang="ko"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>배터리 품질 분석 대시보드</title><style>
*{{box-sizing:border-box}}body{{margin:0;background:#f4f7fb;color:#1f2937;font-family:"Malgun Gothic",sans-serif}}
header{{padding:28px 5%;background:#172554;color:white}}header h1{{margin:0 0 10px}}header p{{color:#dbeafe}}
main{{width:min(1400px,94%);margin:24px auto 50px}}.kpis{{display:grid;grid-template-columns:repeat(auto-fit,minmax(160px,1fr));gap:16px}}
.kpi,.panel{{background:white;border:1px solid #e2e8f0;border-radius:14px;padding:20px;margin-bottom:20px;overflow-x:auto}}
.kpi .label{{color:#64748b}}.kpi .value{{font-size:26px;font-weight:bold;margin-top:10px}}
.stats-table{{width:100%;border-collapse:collapse;font-size:13px}}.stats-table th,.stats-table td{{padding:10px;border-bottom:1px solid #e2e8f0;text-align:right;white-space:nowrap}}
.stats-table th{{background:#eff6ff}}.stats-table th:first-child,.stats-table td:first-child{{text-align:left}}.note{{font-size:13px;color:#64748b;line-height:1.7}}
</style></head><body><header><h1>배터리 품질 분석 대시보드</h1>
<p>입력 파일: {html.escape(source_name)} | 생성 시간: {datetime.now():%Y-%m-%d %H:%M:%S}</p><p>규격 파일: Raw_Data/{SPEC_FILE.name}</p></header><main>
<div class="kpis"><div class="kpi"><div class="label">전체 데이터 수</div><div class="value">{summary["total"]:,}</div></div>
<div class="kpi"><div class="label">전체 OK</div><div class="value">{ok_text}</div></div><div class="kpi"><div class="label">전체 NG</div><div class="value">{ng_text}</div></div>
<div class="kpi"><div class="label">미판정/미측정 포함</div><div class="value">{unclassified_text}</div></div><div class="kpi"><div class="label">합격률 (OK/(OK+NG))</div><div class="value">{yield_text}</div></div></div>
<section class="panel"><h2>기술통계 요약</h2>{stats_html}</section>
<section class="panel"><h2>항목별 OK/NG 및 결측 요약</h2><p class="note">항목별 합격률과 불량률은 해당 항목에서 OK 또는 NG로 판정된 값만 분모로 사용합니다. 미측정·미판정 값은 비율 계산에서 제외합니다.</p>{item_html}</section>
<section class="panel"><h2>공정별 OK/NG 비교</h2>{process_html}</section>
<section class="panel"><h2>규격 및 공정능력 분석</h2><p class="note">LSL·USL이 모두 있고 표준편차가 양수일 때 Cp/Cpk를 계산합니다. 한쪽 규격만 있으면 Cpu 또는 Cpl만 표시합니다. 공정능력 지수는 공정 안정성, 측정 시스템, 표본 대표성을 확인한 뒤 해석하세요. 이 프로젝트의 데이터와 규격이 가상이라면 실제 제품 판정 기준으로 사용하지 마세요.</p>{cap_html}</section>
{graphs}</main></body></html>"""
    HTML_OUTPUT.write_text(page, encoding="utf-8")
    return HTML_OUTPUT


def save_chart_images(df, specs):
    warnings.filterwarnings("ignore", category=UserWarning)
    plt.rcParams["font.family"] = "Malgun Gothic"
    plt.rcParams["axes.unicode_minus"] = False
    outputs = []
    for col in get_numeric_columns(df):
        vals = df[col].dropna()
        fig, ax = plt.subplots(figsize=(8, 4.5))
        ax.hist(vals, bins="auto", edgecolor="black", alpha=0.8)
        sp = specs.get(col, {})
        for field, label, color in (("lsl", "LSL", "red"), ("target", "Target", "blue"), ("usl", "USL", "red")):
            if sp.get(field) is not None:
                ax.axvline(sp[field], linestyle="--", color=color, label=f"{label}={sp[field]:g}")
        if any(sp.get(k) is not None for k in ("lsl", "target", "usl")): ax.legend()
        ax.set_title(f"{DISPLAY_NAMES[col]} 분포"); ax.set_xlabel(DISPLAY_NAMES[col]); ax.set_ylabel("데이터 수")
        ax.grid(axis="y", alpha=.25)
        path = FIGURE_DIR / f"{col}_distribution.png"
        fig.tight_layout(); fig.savefig(path, dpi=150); plt.close(fig)
        outputs.append((f"{DISPLAY_NAMES[col]} 분포", path))

        jcol = f"{col}_판정"
        if jcol in df.columns:
            groups, names = [], []
            for label in ("OK", "NG"):
                group = df.loc[df[jcol] == label, col].dropna()
                if not group.empty:
                    groups.append(group.values); names.append(label)
            if groups:
                fig, ax = plt.subplots(figsize=(8, 4.5))
                ax.boxplot(groups, tick_labels=names, showfliers=True)
                ax.set_title(f"{DISPLAY_NAMES[col]} 항목 자체 OK/NG 비교"); ax.set_xlabel("해당 항목 판정"); ax.set_ylabel(DISPLAY_NAMES[col])
                ax.grid(axis="y", alpha=.25)
                path = FIGURE_DIR / f"{col}_item_ok_ng.png"
                fig.tight_layout(); fig.savefig(path, dpi=150); plt.close(fig)
                outputs.append((f"{DISPLAY_NAMES[col]} 자체 OK/NG 비교", path))
    cols = get_numeric_columns(df)
    if len(cols) >= 2:
        corr = df[cols].corr()
        fig, ax = plt.subplots(figsize=(8, 5))
        im = ax.imshow(corr.values, vmin=-1, vmax=1, cmap="coolwarm")
        names = [DISPLAY_NAMES[c] for c in corr.columns]
        ax.set_xticks(range(len(names))); ax.set_yticks(range(len(names)))
        ax.set_xticklabels(names, rotation=30, ha="right"); ax.set_yticklabels(names)
        for r in range(len(names)):
            for c in range(len(names)): ax.text(c, r, f"{corr.values[r,c]:.2f}", ha="center", va="center")
        ax.set_title("품질 항목 간 상관관계"); fig.colorbar(im, ax=ax, label="상관계수")
        path = FIGURE_DIR / "correlation_heatmap.png"
        fig.tight_layout(); fig.savefig(path, dpi=150); plt.close(fig)
        outputs.append(("품질 항목 간 상관관계", path))
    return outputs


def add_title(slide, title):
    if slide.shapes.title:
        slide.shapes.title.text = title
        for para in slide.shapes.title.text_frame.paragraphs:
            for run in para.runs:
                run.font.name = "Malgun Gothic"; run.font.size = Pt(24); run.font.bold = True


def add_text(slide, text, left, top, width, height, size=16):
    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    shape.text_frame.word_wrap = True; shape.text_frame.text = text
    for para in shape.text_frame.paragraphs:
        for run in para.runs:
            run.font.name = "Malgun Gothic"; run.font.size = Pt(size)
    return shape


def build_powerpoint(df, stats, summary, source_name, specs, capability, item_summary, process_summary):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "배터리 품질 분석 보고서"
    if len(slide.placeholders) > 1:
        slide.placeholders[1].text = f"입력 파일: {source_name}\n생성 시간: {datetime.now():%Y-%m-%d %H:%M}\n통계 및 규격 분석 결과"
    add_title(slide, "배터리 품질 분석 보고서")

    slide = prs.slides.add_slide(prs.slide_layouts[5]); add_title(slide, "1. 분석 개요")
    y = "N/A" if summary["yield_rate"] is None else f'{summary["yield_rate"]:.2f}%'
    add_text(slide, f"입력 파일: {source_name}\n전체 데이터 수: {summary['total']:,}\n전체 OK: {summary['ok']}\n전체 NG: {summary['ng']}\n미판정/미측정: {summary['unclassified']}\n합격률(OK/(OK+NG)): {y}\n규격 파일: Raw_Data/{SPEC_FILE.name}", .8, 1.6, 11.7, 4.8, 19)

    slide = prs.slides.add_slide(prs.slide_layouts[5]); add_title(slide, "2. 기술통계 요약")
    tdf = stats.reset_index().rename(columns={"index": "품질 항목"})
    table = slide.shapes.add_table(len(tdf)+1, len(tdf.columns), Inches(.35), Inches(1.5), Inches(12.6), Inches(min(4.8, .45*(len(tdf)+1)))).table
    for c, name in enumerate(tdf.columns): table.cell(0,c).text = str(name)
    for r, row in enumerate(tdf.itertuples(index=False, name=None), 1):
        for c, val in enumerate(row): table.cell(r,c).text = "—" if pd.isna(val) else (f"{val:.4f}" if isinstance(val,(float,np.floating)) else str(val))
    for row in table.rows:
        for cell in row.cells:
            for para in cell.text_frame.paragraphs:
                for run in para.runs: run.font.name="Malgun Gothic"; run.font.size=Pt(9)

    slide = prs.slides.add_slide(prs.slide_layouts[5]); add_title(slide, "3. 항목별 OK/NG 요약")
    view = item_summary.copy()
    for col in ("합격률(OK/(OK+NG))", "불량률(NG/(OK+NG))"):
        view[col] = view[col].map(lambda x: "—" if pd.isna(x) else f"{x:.2f}%")
    table = slide.shapes.add_table(len(view)+1, len(view.columns), Inches(.35), Inches(1.5), Inches(12.6), Inches(min(4.8, .55*(len(view)+1)))).table
    for c, name in enumerate(view.columns): table.cell(0,c).text = str(name)
    for r, row in enumerate(view.itertuples(index=False, name=None), 1):
        for c, val in enumerate(row): table.cell(r,c).text = str(val)
    for row in table.rows:
        for cell in row.cells:
            for para in cell.text_frame.paragraphs:
                for run in para.runs: run.font.name="Malgun Gothic"; run.font.size=Pt(8)

    if not process_summary.empty:
        slide = prs.slides.add_slide(prs.slide_layouts[5]); add_title(slide, "4. 공정별 OK/NG 비교")
        pview = process_summary.copy()
        pview["합격률(OK/(OK+NG))"] = pview["합격률(OK/(OK+NG))"].map(lambda x: "—" if pd.isna(x) else f"{x:.2f}%")
        table = slide.shapes.add_table(len(pview)+1, len(pview.columns), Inches(.35), Inches(1.5), Inches(12.6), Inches(min(4.8,.55*(len(pview)+1)))).table
        for c, name in enumerate(pview.columns): table.cell(0,c).text = str(name)
        for r, row in enumerate(pview.itertuples(index=False, name=None), 1):
            for c, val in enumerate(row): table.cell(r,c).text = str(val)
        for row in table.rows:
            for cell in row.cells:
                for para in cell.text_frame.paragraphs:
                    for run in para.runs: run.font.name="Malgun Gothic"; run.font.size=Pt(9)

    slide = prs.slides.add_slide(prs.slide_layouts[5]); add_title(slide, "4. 규격 및 공정능력 분석")
    cols = ["품질 항목","데이터 수","Target","LSL","USL","평균","표준편차","평균-Target","Cp","Cpu","Cpl","Cpk"]
    view = capability[cols]
    table = slide.shapes.add_table(len(view)+1, len(cols), Inches(.25), Inches(1.55), Inches(12.85), Inches(min(4.8,.5*(len(view)+1)))).table
    for c, name in enumerate(cols): table.cell(0,c).text = name
    for r, row in enumerate(view.itertuples(index=False, name=None), 1):
        for c, val in enumerate(row): table.cell(r,c).text = "—" if val is None or pd.isna(val) else (f"{val:.3f}" if isinstance(val,(float,np.floating)) else str(val))
    for row in table.rows:
        for cell in row.cells:
            for para in cell.text_frame.paragraphs:
                for run in para.runs: run.font.name="Malgun Gothic"; run.font.size=Pt(8)

    for title, image_path in save_chart_images(df, specs):
        slide = prs.slides.add_slide(prs.slide_layouts[5]); add_title(slide, title)
        slide.shapes.add_picture(str(image_path), Inches(1), Inches(1.5), width=Inches(11.3))

    slide = prs.slides.add_slide(prs.slide_layouts[5]); add_title(slide, "분석 결과 해석 시 유의사항")
    notes = ["상관관계는 인과관계를 의미하지 않습니다.", "관리한계는 통계적 관리 기준이며 제품 규격과 다릅니다.",
             "Cp/Cpk는 공정 안정성, 측정 시스템, 표본 대표성을 확인한 뒤 해석해야 합니다.",
             "결측값은 미측정으로 표시하며 항목별 합격률 분모에서 제외합니다.",
             "공정 비교를 하려면 입력 데이터에 공정명/공정 컬럼이 필요합니다.",
             "가상 데이터 및 연습용 규격은 실제 제품의 출하 판정 기준이 아닙니다."]
    add_text(slide, "\n\n".join("• "+n for n in notes), .8, 1.6, 11.7, 5.2, 15)
    prs.save(PPT_OUTPUT)
    return PPT_OUTPUT


def run_analysis():
    initialize_log(); write_log("분석 작업을 시작합니다.")
    source = find_input_file(); write_log(f"입력 데이터 파일: {source.name}")
    raw = read_input_file(source); write_log(f"원본 데이터 크기: {raw.shape}")
    df = prepare_data(raw); write_log(f"전처리 후 데이터 크기: {df.shape}")
    specs = load_specifications(df)
    df = apply_spec_judgments(df, specs)
    df.to_csv(RESULT_CSV, index=False, encoding="utf-8-sig")
    write_log(f"판정 결과 CSV 생성: {RESULT_CSV}")
    write_log(f"규격 설정 파일: {SPEC_FILE}")
    stats = calculate_statistics(df)
    summary = calculate_quality_summary(df)
    item_summary = calculate_item_summary(df)
    process_summary = calculate_process_summary(df)
    capability = calculate_capability(df, specs)
    html_path = build_dashboard(df, stats, summary, source.name, specs, capability, item_summary, process_summary)
    write_log(f"HTML 대시보드 생성: {html_path}")
    ppt_path = build_powerpoint(df, stats, summary, source.name, specs, capability, item_summary, process_summary)
    write_log(f"PowerPoint 보고서 생성: {ppt_path}")
    write_log(f"전체 판정 요약: OK={summary['ok']}, NG={summary['ng']}, 미판정/미측정={summary['unclassified']}, 합격률={summary['yield_rate']}")
    write_log("모든 분석 작업이 완료되었습니다.")
    return html_path, ppt_path


def main():
    try:
        html_path, _ = run_analysis()
        if webview is not None:
            webview.create_window("배터리 품질 분석 자동화", html_path.resolve().as_uri(), width=1400, height=900, resizable=True)
            webview.start()
        else:
            webbrowser.open(html_path.resolve().as_uri())
            print(f"대시보드를 브라우저에서 열었습니다: {html_path}")
    except Exception as error:
        message = f"프로그램 실행 중 오류가 발생했습니다.\n\n{type(error).__name__}: {error}\n\n로그 파일: {LOG_OUTPUT}"
        try: write_log(message); write_log(traceback.format_exc())
        except Exception: pass
        print("\n"+message+"\n상세 오류:"); traceback.print_exc()
        try: input("\nEnter 키를 누르면 종료합니다...")
        except EOFError: pass


if __name__ == "__main__":
    main()
